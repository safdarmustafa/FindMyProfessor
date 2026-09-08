from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.cv.parsers.base import CvParser, ParseError


class DocxParser(CvParser):
    format_id = "docx"

    def extract_text(self, content: bytes) -> str:
        try:
            document = Document(BytesIO(content))
        except (PackageNotFoundError, ValueError, KeyError) as exc:
            raise ParseError("The Word document could not be opened.") from exc
        parts: list[str] = []
        for paragraph in document.paragraphs:
            parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
